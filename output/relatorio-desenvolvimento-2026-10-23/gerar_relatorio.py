from pathlib import Path
import json, re, subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
NAME = 'relatorio_desenvolvimento_o_segredo_de_varginha_23-10-2026'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.left_margin = Cm(3), Cm(3)
sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(2), Cm(1.25)

def set_font(st, size=12, bold=False):
    st.font.name='Times New Roman'; st.font.size=Pt(size)
    st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
        st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)

normal=doc.styles['Normal']; set_font(normal)
pf=normal.paragraph_format
pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
pf.first_line_indent=Cm(1.25); pf.line_spacing=1.5
pf.space_after=Pt(0); pf.widow_control=True
for name,size in [('Title',16),('Heading 1',16),('Heading 2',12),('Heading 3',12)]:
    st=doc.styles[name]; set_font(st,size,True)
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before=Pt(18 if name=='Heading 2' else 0)
    st.paragraph_format.space_after=Pt(18)
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
for name in ['Caption','Source','Reference','Front','ListEntry','TOC Heading']:
    if name not in doc.styles:
        doc.styles.add_style(name,1)
    st=doc.styles[name];set_font(st,10 if name in ['Caption','Source'] else 12)
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.line_spacing=1.0
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Caption'].paragraph_format.keep_with_next=True
doc.styles['Source'].paragraph_format.space_after=Pt(12)
doc.styles['Reference'].paragraph_format.space_after=Pt(12)
doc.styles['Reference'].paragraph_format.keep_together=True
doc.styles['Front'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['TOC 1','TOC 2']:
    if name not in doc.styles: doc.styles.add_style(name,1)
    st=doc.styles[name];set_font(st,12,name=='TOC 1')
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.space_before=Pt(0)
    st.paragraph_format.space_after=Pt(0)

def para(text='', style=None):
    return doc.add_paragraph(text,style)

def front(text, before=0, after=0, bold=False, size=12):
    p=para(text,'Front');p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    for r in p.runs:r.bold=bold;r.font.size=Pt(size)
    return p

def page():doc.add_page_break()

def field(p, code, placeholder=''):
    run=p.add_run(); b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');run._r.append(b)
    r=p.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' '+code+' ';r._r.append(i)
    r=p.add_run();s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');r._r.append(s)
    p.add_run(placeholder)
    r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)

def section(start=None):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.header.is_linked_to_previous=False
    s.footer.is_linked_to_previous=False
    for n in s._sectPr.findall(qn('w:pgNumType')):
        s._sectPr.remove(n)
    if start is not None:
        n=OxmlElement('w:pgNumType');n.set(qn('w:start'),str(start));s._sectPr.append(n)
    return s

def unnumbered(title, include=False, newpage=True):
    if newpage:page()
    p=para(title,'Heading 1' if include else 'Title')
    p.paragraph_format.page_break_before=False
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p

front('ESCOLA ESTADUAL DEPUTADO DOMINGOS DE FIGUEIREDO',bold=True)
front('CURSO TÉCNICO EM DESENVOLVIMENTO DE SISTEMAS',before=12)
front('TURMA DO 3º ANO DE DESENVOLVIMENTO DE SISTEMAS',before=72)
p=front('O SEGREDO DE VARGINHA',before=105,bold=True,size=16);p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
front('RELATÓRIO DO DESENVOLVIMENTO DO PROJETO',before=12,bold=True)
front('Dos primeiros registros ao estágio de 23 de outubro de 2026',before=12)
front('Varginha',before=175)
front('2026',before=8)

section(1)
front('TURMA DO 3º ANO DE DESENVOLVIMENTO DE SISTEMAS')
p=front('O SEGREDO DE VARGINHA',before=100,bold=True,size=16);p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
front('RELATÓRIO DO DESENVOLVIMENTO DO PROJETO',before=12,bold=True)
p=para('Relatório de desenvolvimento apresentado à Escola Estadual Deputado Domingos de Figueiredo, no 3º ano do curso técnico em Desenvolvimento de Sistemas, para registro e análise da trajetória do projeto O Segredo de Varginha.','Front')
p.paragraph_format.left_indent=Cm(8);p.paragraph_format.space_before=Pt(90);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
p=para('Professores: Renan e Alex.','Front');p.paragraph_format.left_indent=Cm(8);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
front('Varginha',before=130)
front('2026',before=8)

unnumbered('RESUMO')
para('Este relatório apresenta o desenvolvimento do jogo O Segredo de Varginha, produzido no contexto do curso técnico em Desenvolvimento de Sistemas da Escola Estadual Deputado Domingos de Figueiredo, em Varginha. O objetivo é reconstruir a trajetória do projeto e analisar seus aspectos narrativos, técnicos, visuais, sonoros e organizacionais até 23 de outubro de 2026. Adotou-se pesquisa documental sobre o histórico Git, documentos de design, código, capturas e registros de qualidade. O primeiro commit acessível é de 11 de setembro e já contém um protótipo; a data da concepção anterior não pôde ser determinada. O desenvolvimento ampliou a campanha e depois condensou seu percurso em 14 capítulos e cinco atos, preservando investigação, enigmas, aliados, combate e uma conclusão implementada. Foram integrados sistemas de persistência, controles, interface, animação, efeitos e ambientação sonora. A entrega local mais recente identificada é a 0.4.8-beta, compilada em Unity 6000.6.0f1, com zero erros e um aviso. A revisão de falas dessa versão possui 29 testes EditMode e três casos no Player aprovados em seu escopo. Esses resultados não equivalem a uma validação integral da campanha em uma única execução. Permanecem pendentes a dublagem humana do lote aprovado, testes externos de experiência, medição de desempenho em outro computador e consolidação de documentos e créditos. Conclui-se que o projeto alcançou uma beta local com percurso narrativo completo em sua implementação, necessitando de estabilização e avaliação humana para fundamentar a próxima distribuição.')
p=para('Palavras-chave: desenvolvimento de jogos; investigação; Unity; pixel art; Varginha.');p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(18)

unnumbered('ABSTRACT')
para('This report presents the development of O Segredo de Varginha, a game produced in the technical Systems Development course at Escola Estadual Deputado Domingos de Figueiredo, in Varginha, Brazil. It reconstructs the project history and examines narrative, technical, visual, audio and organizational aspects up to October 23, 2026. The method is documentary analysis of Git history, design documents, source code, screenshots and quality records. The earliest available commit dates from September 11 and already contains a prototype; the earlier conception date could not be established. Development expanded the campaign before condensing it into fourteen chapters and five acts, retaining investigation, puzzles, allies, combat and one implemented ending. Persistence, controls, interface, animation, sound effects and ambience were integrated. The latest identified local release is version 0.4.8-beta, built with Unity 6000.6.0f1 with no errors and one warning. Its dialogue revision has twenty-nine passing targeted EditMode tests and three passing Windows Player cases. These results do not establish a complete campaign validation in a single execution. Human voice recording, external player evaluation, performance measurements on another computer and consolidation of documentation and credits remain pending. The project has reached a local beta with an implemented narrative route through its ending, while further stabilization and human evaluation are required to support its next distribution.')
p=para('Keywords: game development; investigation; Unity; pixel art; Varginha.');p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(18)

figure_titles=['Escola Industrial no acervo da remasterização','Diálogo escolar registrado na revisão de 23 de outubro','Mesa do café e apresentação da investigação da chave']
quadro_titles=['Campanha vigente e identificadores preservados','Riscos e medidas recomendadas','Situação dos principais aspectos do projeto']
table_titles=['Dimensão estática do acervo consultado','Verificações selecionadas das versões recentes']

for title,kind,items in [('LISTA DE ILUSTRAÇÕES','Figura',figure_titles),('LISTA DE QUADROS','Quadro',quadro_titles),('LISTA DE TABELAS','Tabela',table_titles)]:
    unnumbered(title)
    for n,t in enumerate(items,1):
        p=para(f'{kind} {n} – {t}','ListEntry')
        p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(12)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9),WD_ALIGN_PARAGRAPH.RIGHT,1)
        p.add_run('\t');field(p,f'PAGEREF {kind}{n} \\h','0')

unnumbered('LISTA DE ABREVIATURAS E SIGLAS')
for sig,desc in [('ABNT','Associação Brasileira de Normas Técnicas'),('C#','Linguagem de programação C Sharp'),('GDD','Game Design Document ou documento de design de jogo'),('HUD','Heads-up display ou interface de informações do jogo'),('IA','Inteligência artificial'),('ID','Identificador'),('JSON','JavaScript Object Notation'),('LDD','Level Design Document ou documento de design de níveis'),('MDA','Mechanics Dynamics Aesthetics'),('PCM','Pulse-code modulation'),('QA','Quality assurance ou garantia da qualidade'),('UI','User interface ou interface do usuário'),('URP','Universal Render Pipeline'),('VFX','Visual effects ou efeitos visuais'),('WAV','Formato de arquivo de áudio Waveform'),('XML','Extensible Markup Language')]:
    p=para(f'{sig}\t{desc}','ListEntry');p.paragraph_format.line_spacing=1.5;p.paragraph_format.tab_stops.add_tab_stop(Cm(2.1))

unnumbered('SUMÁRIO')
field(para('', 'ListEntry'),'TOC \\o "1-2" \\h \\z \\u','Atualizar sumário')

s=section()
p=s.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Cm(0)
field(p,'PAGE','0')
for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)

book_id=0;fig_no=0
def caption(kind,n,title):
    global book_id
    p=para(f'{kind} {n} – {title}','Caption')
    book_id+=1
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(book_id));a.set(qn('w:name'),f'{kind}{n}')
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(book_id))
    p._p.insert(0,a);p._p.append(b)
    return p

def source(text):para('Fonte: '+text,'Source')

def table(kind,n,title,headers,rows,widths,src):
    caption(kind,n,title)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,w,h in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=h
    trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,w,text in zip(cells,widths,row):c.width=Cm(w);c.text=str(text)
    for rn,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        for cn,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side)
                draw=(kind=='Quadro') or (side=='top' and rn==0) or (side=='bottom' and (rn==0 or rn==len(t.rows)-1))
                e.set(qn('w:val'),'single' if draw else 'nil');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9' if kind=='Quadro' else '000000');borders.append(e)
            tcpr.append(borders)
            if rn==0 and kind=='Quadro':
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=False
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cn>0 or len(headers)==2 else WD_ALIGN_PARAGRAPH.CENTER
                if kind=='Tabela' and n==1 and cn==1:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=rn==0
    source(src)

def campaign():
    rows=[
    ('1','O Caso de Varginha','Infância, brilho, explosão e memória','1'),
    ('2','A Caixa Esquecida','Rotina, chave, caixa e primeiro fragmento','2'),
    ('3','Entre Aulas e Pistas','Fotografia, vulto, horário e Renan','4'),
    ('4','Fragmentos','Documentos e montagem do mapa','6'),
    ('5','A Mata','Orientação, perseguição e Fábio','7'),
    ('6','A Âncora','Livro do Tombo e registro de Edelzio','8'),
    ('7','Marcas no Fusca','Comparação com Ouzana e reagente','9 e 10'),
    ('8','Arquivos Alterados','Comparação de papel e arquivo digital','11'),
    ('9','O Casarão de 1898','Planta, passagem, chave e porão','12 e 13'),
    ('10','O Que Fábio Escondeu','Provas e página ocultada','14'),
    ('11','Antes da Explosão','Lembrança da casa em 1996','15'),
    ('12','A Criatura e o Acordo','Evidências e reconstrução do acordo','18'),
    ('13','A Batalha da Manifestação','Preparação, combate e calibração','20'),
    ('14','O Retorno','Travessia, encerramento e escola','21')]
    table('Quadro',1,quadro_titles[0],['Cap.','Título','Função principal','IDs'],rows,[1.1,4.7,8.1,2.1],'elaboração para este relatório a partir de Projeto O Segredo de Varginha (2026b, 2026c). As áreas complementares não estão individualizadas nesta síntese.')

def inventory():
    table('Tabela',1,table_titles[0],['Recorte','Quantidade'],[
    ('Commits alcançáveis até o último commit',49),('Arquivos C# em Assets/Scripts/Game',183),('Arquivos C# em Assets/Tests/EditMode',32),('Arquivos C# em Assets/Tests/PlayMode',46),('Cenas em Assets/Scenes',30),('Cenas incluídas na build 0.4.8-beta',19),('Capítulos apresentados na campanha',14)], [12.5,3.5], 'levantamento documental do repositório e manifesto da build de 23 out. 2026 (Projeto O Segredo de Varginha, 2026c, 2026e, 2026k). Arquivos não equivalem a casos de teste.')

def tests():
    rows=[('21/10','0.4.1','348/348','270/270¹','4/4²'),('21/10','0.4.2','367/367','—','5/5'),('21/10','0.4.3','378/378','—','8/8'),('21/10','0.4.4','389/389','—','10/10³'),('22/10','0.4.5','398/398','—','16/16⁴'),('22/10','0.4.6','398/398','—','3/3'),('23/10','0.4.8','29/29⁵','—','3/3⁵')]
    table('Tabela',2,table_titles[1],['Data','Beta','EditMode','PlayMode','Player'],rows,[2,2.5,4,4,3.5],'Projeto O Segredo de Varginha (2026i, 2026j, 2026k), resumos das entregas. Valores aprovados/total, não cumulativos.')
    p=para('Notas: ¹ Último resultado por caso, sem rodada integral limpa. ² Duas execuções. ³ Quatro repetições adicionais não somadas. ⁴ Casos únicos consolidados. ⁵ Recorte de falas; houve rodada inicial de 2/3 no Player. O travessão indica resultado não atribuído ao recorte.','Source')
    p.paragraph_format.keep_together=True

def risks():
    table('Quadro',2,quadro_titles[1],['Risco','Consequência','Medida recomendada'],[
    ('Documentos em versões diferentes','Apresentação de escopo ou versão incorretos','Indicar documento vigente e manter os anteriores como histórico.'),
    ('Migração de saves','Perda de progresso ou bloqueio de sequência','Repetir casos novos, antigos e de retorno na candidata.'),
    ('Regressão visual e de interação','Pistas encobertas ou aproximação inconsistente','Conferir profundidade, suporte físico e fluxo em contexto.'),
    ('Validação humana limitada','Dificuldade e clareza pouco conhecidas','Aplicar partida completa e teste com iniciantes.'),
    ('Permissões e créditos incompletos','Dificuldade de comprovar condições de publicação','Consolidar origem e autorizações por recurso.'),
    ('Dublagem ainda pendente','Descompasso entre roteiro e áudio recebido','Usar lote aprovado e mapa de IDs; validar a integração.')], [3.5,5.2,7.3], 'análise deste relatório com base nas limitações registradas no acervo.')

def status():
    table('Quadro',3,quadro_titles[2],['Aspecto','Situação verificada','Limite ou próxima ação'],[
    ('Narrativa e percurso','14 capítulos, cinco atos, uma conclusão implementada','Finais alternativos continuam fora do escopo disponível.'),
    ('Investigação e puzzles','Regras, evidências, caderno e dicas integrados','Avaliar compreensão por jogadores iniciantes.'),
    ('Combate e apoios','Confronto final, aliados e perfis de apoio documentados','Avaliar equilíbrio em partida completa.'),
    ('Arte e animação','Recursos originais preservados e produção complementar integrada','Continuar revisão contextual de oclusão e escala.'),
    ('Som e voz','Efeitos e ambientes integrados; roteiro definitivo preparado','Receber e integrar dublagem; resolver ED03.'),
    ('Testes','Resultados por versão e recortes preservados','Executar regressão da candidata e avaliação externa.'),
    ('Build','0.4.8-beta local gerada em Windows, zero erros e um aviso','Conferir pacote e publicação antes da divulgação.'),
    ('Documentação','Histórico amplo e relatórios recentes disponíveis','Consolidar entrada, créditos e versão apresentada.')], [3.0,6.4,6.6], 'síntese deste relatório a partir de Projeto O Segredo de Varginha (2026b, 2026f, 2026i, 2026j, 2026k).')

content=(OUT/'conteudo.md').read_text(encoding='utf-8')
chunks=re.split(r'\n\s*\n',content.strip())
for chunk in chunks:
    chunk=chunk.strip()
    if chunk.startswith('# '):
        number,title=chunk[2:].split(' ',1)
        p=para(number+'\t'+title,'Heading 1')
        p.paragraph_format.left_indent=Cm(0.8)
        p.paragraph_format.first_line_indent=Cm(-0.8)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))
        if chunk.startswith('# 1 '):p.paragraph_format.page_break_before=False
    elif chunk.startswith('## '):
        p=para(chunk[3:],'Heading 2')
        if chunk.startswith('## 11.4 '):p.paragraph_format.page_break_before=True
    elif chunk.startswith('@fig|'):
        _,title,path,src=chunk.split('|')
        fig_no+=1;caption('Figura',fig_no,title)
        p=para();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/path),width=Cm(16))
        source(src)
    elif chunk in ['@campaign','@inventory','@tests','@risks','@status']:
        globals()[chunk[1:]]()
    else:para(chunk.replace('\n',' '))

unnumbered('REFERÊNCIAS',include=True)
refs=[
'ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 14724: informação e documentação — trabalhos acadêmicos — apresentação. 4. ed. Rio de Janeiro: ABNT, 2024. Disponível em: https://www.ccs.ufpb.br/coordfarm/wp-content/uploads/sites/21/sites/454/2026/03/ABNT_NBR_14724_2024-1.pdf. Acesso em: 23 out. 2026.',
'CHACON, Scott; STRAUB, Ben. Pro Git. 2. ed. [S. l.]: Apress, 2014. Seção About Version Control. Disponível em: https://git-scm.com/book/en/v2/Getting-Started-About-Version-Control. Acesso em: 23 out. 2026.',
'HUNICKE, Robin; LEBLANC, Marc; ZUBEK, Robert. MDA: a formal approach to game design and game research. In: AAAI WORKSHOP ON CHALLENGES IN GAME AI, 2004. Anais [...]. [S. l.]: AAAI, 2004. Disponível em: https://www.cs.northwestern.edu/~hunicke/MDA.pdf. Acesso em: 23 out. 2026.',
'MODELO ABTN. [S. l.: s. n.], [s. d.]. 10 p. Arquivo PDF fornecido pelo solicitante como referência de apresentação.',
'PROJETO O SEGREDO DE VARGINHA. Arquitetura do protótipo. Varginha, 2026a. Documento interno. Localização: Docs/ARCHITECTURE.md. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Campanha experimental com 14 fases em 5 atos. Varginha, 2026b. Documento interno, com atualização da abertura em 20 out. Localização: Docs/CampanhaOficial/CAMPANHA_14_FASES.md. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Código e configuração do projeto Unity. Varginha, 2026c. Repositório local, incluindo Assets/Scripts, Packages e ProjectSettings. Estado consultado em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Documentos de design de jogo e de níveis. Varginha, 2026d. Conjunto documental interno. Localização: Docs/CampanhaOficial/Referencias/GDD_Usuario.md e O_Segredo_de_Varginha_LDD.md; versão consolidada histórica em Docs/GDD.md. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Histórico Git. Varginha, 2026e. 49 commits alcançáveis entre 8dd459a, de 11 set. 2026, e a2946be, de 22 out. 2026. Consulta local em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Produção visual sonora e auditoria de apoios. Varginha, 2026f. Conjunto documental interno: ASSETS_IA.md; INTEGRACAO_AUDIO_2026-10-21.md; CREDITOS_AUDIO_RECEBIDO_2026-10-21.md; CREDITOS_AUDIO_EDELZIO_2026-10-22.md; AUDITORIA_APOIOS_2026-10-21.md. Diretório: Docs/CampanhaOficial. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Relatórios de campanha remasterização alpha e beta. Varginha, 2026g. Conjunto documental interno. Diretórios: Docs/CampanhaOficial; Docs/Remaster20261008; Docs/QAAlpha01; Docs/Alpha03; Docs/Beta04. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Roteiros e mapas de dublagem de Edelzio. Varginha, 2026h. Conjunto de materiais internos, incluindo roteiro_edelzio_definitivo.pdf e mapas de voz de 23 out. Diretórios: output/pdf e output/dublagem-2026-10-23. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Verificações técnicas de 20 e 21 de outubro. Varginha, 2026i. Relatórios, JSONs e resultados de testes das entregas 0.4.1 a 0.4.4-beta e da investigação da chave. Diretórios discriminados no Apêndice B. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Verificações técnicas de 22 de outubro. Varginha, 2026j. Relatórios, capturas e manifestos de interações da casa, mesa escolar e áudio. Diretórios: output/house-interactions-2026-10-22; output/school-desk-2026-10-22; output/audio-edelzio-2026-10-22. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Verificações técnicas da revisão de falas de 23 de outubro. Varginha, 2026k. Relatório, resultados EditMode e Player e manifesto da build 0.4.8-beta. Localização: output/dialogue-casual-2026-10-23. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Website de O Segredo de Varginha. Varginha, 2026l. Código estático e documentação interna. Localização: Website/README.md e Website/dist. Consulta em: 23 out. 2026.',
'UNITY TECHNOLOGIES. Play mode tests. [S. l.]: Unity Technologies, [s. d.]. Unity User Manual. Disponível em: https://docs.unity.com/en-us/engine/6000.3/manual/programming-environment/test-framework-introduction/overview/general-introduction/play-mode-tests. Acesso em: 23 out. 2026.',
'UNIVERSIDADE FEDERAL DO ESPÍRITO SANTO. Sistema Integrado de Bibliotecas. Normalização. Vitória: UFES, 2026. Atualizado em 10 jul. 2026. Disponível em: https://biblioteca.ufes.br/normalizacao. Acesso em: 23 out. 2026.'
]
emphases=['ABNT NBR 14724','Pro Git','MDA: a formal approach to game design and game research','MODELO ABTN','Arquitetura do protótipo','Campanha experimental com 14 fases em 5 atos','Código e configuração do projeto Unity','Documentos de design de jogo e de níveis','Histórico Git','Produção visual sonora e auditoria de apoios','Relatórios de campanha remasterização alpha e beta','Roteiros e mapas de dublagem de Edelzio','Verificações técnicas de 20 e 21 de outubro','Verificações técnicas de 22 de outubro','Verificações técnicas da revisão de falas de 23 de outubro','Website de O Segredo de Varginha','Play mode tests','Normalização']
for ref,emphasis in zip(refs,emphases):
    p=para('', 'Reference');before,after=ref.split(emphasis,1)
    p.add_run(before);p.add_run(emphasis).bold=True;p.add_run(after)

unnumbered('APÊNDICE A — MARCOS DO HISTÓRICO',include=True)
para('A relação abaixo seleciona marcos dos 49 commits alcançáveis. As descrições foram sintetizadas para este relatório e não são transcrições literais das mensagens. Datas e hashes permitem localizar os estados. A última linha registra alteração local, sem atribuir-lhe um commit inexistente.')
milestones=[
('11/09','8dd459a','Protótipo com jogabilidade e pixel art.'),('12/09','ca185c2','Testes de fases e aliados.'),('14/09','c331760','Comandos, apresentação e cenários.'),('15/09','eac5a90','Guia do prólogo e qualidade.'),('17/09','9e53e4a','Derrota, aliados e testes.'),('22/09','297496b','Sprites e correções.'),('25/09','1fa076c','Combate Punch V2 e testes.'),('25/09','e6debb6','Seleção de lanterna e persistência.'),('28/09','842fd59','Refatoração.'),('29/09','2001d7e','Resgate, iluminação e arte da diocese.'),('30/09','f34e0db','Sala da segunda fase e animação sentada.'),('01/10','4e5248e','Cópia experimental e plano de 20 fases.'),('01/10','a29aabf','Abertura e criança jogável.'),('02/10','1c1d0ad','Mapas e etapas intermediárias.'),('03/10','e97aa4c','Campanha ilustrada e investigação escolar.'),('04/10','ecabd92','Etapas finais e confronto conclusivo.'),('05/10','479328d','Apresentação e fluidez do chefe.'),('07/10','97c8ef9','Condensação de percurso e usabilidade.'),('07/10','412cd5f','Distribuição alpha e site.'),('08/10','e74baa3','Remasterização cinematográfica.'),('09/10','6edf652','Alpha 0.3, mapas, interface e animação.'),('10/10','df2b091','Beta 0.4, laboratório, menu e selo.'),('20/10','44a80af','Investigação da chave e daily.'),('21/10','117d2eb','Abertura adulta e 0.4.1-beta.'),('21/10','6dfd878','Fusca, bifurcações e 0.4.2-beta.'),('21/10','572a296','Prólogo, manhã e 0.4.3-beta.'),('21/10','786143d','Áudio, apoios e 0.4.4-beta.'),('22/10','eb2f8ea','Rotina, cinema e 0.4.5-beta.'),('22/10','4362b53','Janelas, mesa e 0.4.6-beta.'),('22/10','a2946be','Áudio de efeitos e 0.4.7-beta.'),('23/10','Local','Falas aprovadas e build 0.4.8-beta.')]
for dt,sha,desc in milestones:
    p=para(f'{dt}/2026 — {sha} — {desc}','Reference');p.paragraph_format.space_after=Pt(3)
source('elaboração deste relatório a partir do histórico Git e dos relatórios de entrega. O histórico acessível não documenta a concepção anterior a setembro.')

unnumbered('APÊNDICE B — RASTREABILIDADE DAS FONTES',include=True)
para('Os caminhos abaixo são relativos à raiz do projeto consultado. Preservam a localização dos registros sem reproduzir caminhos pessoais do computador. Documentos históricos devem ser lidos em conjunto com os mais recentes. Os conjuntos agrupados nas referências têm finalidade de rastreabilidade, e não indicam um único arquivo com esse título.')
catalog=[
('Planejamento original','Docs/CampanhaOficial/Referencias/GDD_Usuario.md; Docs/CampanhaOficial/Referencias/O_Segredo_de_Varginha_LDD.md'),
('Campanha vigente','Docs/CampanhaOficial/CAMPANHA_14_FASES.md; Assets/Scripts/Game/Varginha/Experiment/CampaignSequence.cs'),
('Versões intermediárias','Docs/CampanhaOficial/PLANO_20_FASES.md; Docs/CampanhaOficial/CAMPANHA_15_FASES.md'),
('Arquitetura e ambiente','Docs/ARCHITECTURE.md; ProjectSettings/ProjectVersion.txt; Packages/manifest.json'),
('Abertura e enigmas','Docs/CampanhaOficial/CHAVE_DESAPARECIDA.md; Docs/CampanhaOficial/PUZZLE_FRAGMENTOS_BIFURCACOES_2026-10-21.md'),
('Prólogo e coerência narrativa','Docs/CampanhaOficial/ROTEIRO_CUTSCENES_GDD_LDD_2026-10-21.md'),
('Histórico de qualidade','Docs/CampanhaOficial/RELATORIO_15_FASES_20261007.md; Docs/QAAlpha01/VALIDACAO.md; Docs/Alpha03/README.md; Docs/Beta04/README.md'),
('Remasterização','Docs/Remaster20261008/RELATORIO_SPRINT_2026-10-09.md; Docs/Remaster20261008/Prints'),
('Investigação da chave','output/key-investigation-2026-10-20/RELATORIO.md'),
('Entrega 0.4.1','output/qa-2026-10-21/RELATORIO.md; output/qa-2026-10-21/summary.json'),
('Entrega 0.4.2','output/cinematic-opening-2026-10-21/RELATORIO.md; output/cinematic-opening-2026-10-21/summary.json'),
('Entrega 0.4.3','output/visual-polish-2026-10-21/RELATORIO.md; output/visual-polish-2026-10-21/summary.json'),
('Entrega 0.4.4','output/audio-integration-2026-10-21/RELATORIO.md; output/audio-integration-2026-10-21/summary.json'),
('Entregas da casa e escola','output/house-interactions-2026-10-22/RELATORIO.md; output/school-desk-2026-10-22/RELATORIO.md'),
('Falas e testes de 23 de outubro','output/dialogue-casual-2026-10-23/RELATORIO.md; editmode-summary.json; player-summary.json; windows-build.json, no mesmo diretório'),
('Rodada inicial preservada','output/dialogue-casual-2026-10-23/player-initial-summary.json; player-initial.xml'),
('Roteiro definitivo','output/pdf/roteiro_edelzio_definitivo.pdf; output/dublagem-2026-10-23/versao-descontraida'),
('Arte áudio e apoios','Docs/CampanhaOficial/ASSETS_IA.md; CREDITOS_AUDIO_RECEBIDO_2026-10-21.md; CREDITOS_AUDIO_EDELZIO_2026-10-22.md; AUDITORIA_APOIOS_2026-10-21.md, no mesmo diretório'),
('Entrada e website','Docs/CampanhaOficial/CONTROLE_XINPUT.md; Website/README.md'),
('Compilação mais recente','Builds/Varginha_0.4.8-beta_2026-10-23/SegredoDeVarginha.exe; manifesto em output/dialogue-casual-2026-10-23/windows-build.json'),
('Ferramentas de evidência','Tools/AgentOutput/README.md; Tools/QA/run_dialogue_feedback.ps1')]
for title,path in catalog:
    p=para(title,'Heading 2');p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(4)
    # Catalog labels are not TOC subsections.
    ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');p._p.get_or_add_pPr().append(ol)
    p=para(path,'Reference')
    for r in p.runs:r.font.size=Pt(10)
    p.paragraph_format.space_after=Pt(6)

unnumbered('APÊNDICE C — PROCEDIMENTO DE CONFERÊNCIA',include=True)
for title,text in [
('Identificação do estado','Registrar a versão do Unity, o hash do commit e as alterações locais. O hash sozinho não representa a build de 23 de outubro, pois há mudanças posteriores ao último commit. Preservar o manifesto e os resultados junto do pacote avaliado.'),
('Consulta ao histórico','Usar git log --reverse para localizar os primeiros registros, git log -1 para identificar o último commit e git status --short para distinguir arquivos locais. Esses comandos de consulta não substituem a inspeção do conteúdo das alterações.'),
('Conferência de resultados','Abrir o resumo da rodada e seu XML correspondente. Conferir total, aprovados, falhas, ignorados, versão, data e método de consolidação. Quando houver falha, consultar o log preservado e registrar se ocorreu nova execução. Não interpretar zero casos como sucesso.'),
('Verificação da candidata','Executar a build identificada, e não uma instância anterior já aberta. Conferir menu, novo jogo, continuidade, casa, escola, pistas, pausa e retorno. Para validar o percurso completo, registrar uma partida do início ao fim e a configuração de hardware.'),
('Proteção de progresso','Separar saves e preferências de teste dos dados de uso pessoal. Antes de testes de migração, preservar cópias e identificar a versão de origem. O runner de 23 de outubro utilizou identificação de produto separada e restauração de configurações ao terminar.'),
('Registro de distribuição','Confirmar que o ZIP corresponde à versão avaliada, que seus arquivos estão íntegros e que o link público entrega esse pacote. Conferir versão no site, apresentação, README e notas da entrega. Uma build local não prova publicação.'),
('Continuidade do relatório','Atualizar a data de corte e acrescentar novas evidências, sem transformar recomendações em tarefas concluídas. Manter acessíveis os resultados anteriores e explicar mudanças de escopo que alterem a interpretação dos indicadores.')]:
    p=para(title,'Heading 2');p.paragraph_format.space_before=Pt(12)
    ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');p._p.get_or_add_pPr().append(ol)
    para(text)

# Prevent automatic heading entries from catalog labels by using a non-outline style.
if 'AppendixLabel' not in doc.styles:
    st=doc.styles.add_style('AppendixLabel',1);st.base_style=doc.styles['Normal'];set_font(st,12,True)
    st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.0
for p in doc.paragraphs:
    if p._p.xpath('./w:pPr/w:outlineLvl[@w:val="9"]'):
        p.style='AppendixLabel'
        p.paragraph_format.space_before=Pt(8)
        p.paragraph_format.space_after=Pt(4)

# Keep title paragraphs and native table text black, including inherited theme settings.
for st in doc.styles:
    if st.type==1:
        st.font.color.rgb=RGBColor(0,0,0)
        # Remove theme font overrides and inherited decorative rules.
        set_font(st,st.font.size.pt if st.font.size else 12, bool(st.font.bold))
        for b in st.element.xpath('./w:pPr/w:pBdr'):
            b.getparent().remove(b)
for p in doc.paragraphs:
    for b in p._p.xpath('./w:pPr/w:pBdr'):
        b.getparent().remove(b)
doc.core_properties.title='O Segredo de Varginha — Relatório do desenvolvimento do projeto'
doc.core_properties.subject='Histórico e análise do desenvolvimento até 23 de outubro de 2026'
doc.core_properties.author='Turma do 3º ano de Desenvolvimento de Sistemas'
doc.core_properties.keywords='Varginha, desenvolvimento de jogos, Unity, relatório, ABNT'
settings=doc.settings.element
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
dest=OUT/(NAME+'.docx');doc.save(dest)
stats={'docx':str(dest),'words_body':len(content.split()),'figures':fig_no,'quadros':len(quadro_titles),'tables':len(table_titles),'primary_sections':13,'references':len(refs),'cutoff':'2026-10-23','format':'A4; Times New Roman 12; 1.5; margins 3/3/2/2 cm'}
(OUT/'autoria-e-escopo.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(stats,ensure_ascii=False))
