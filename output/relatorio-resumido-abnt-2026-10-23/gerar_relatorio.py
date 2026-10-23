from pathlib import Path
import json, re, subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
NAME = 'relatorio_desenvolvimento_varginha_resumido'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.left_margin = Cm(3), Cm(3)
sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(2), Cm(1.25)

def set_font(st, size=12, bold=False):
    st.font.name='Times New Roman'; st.font.size=Pt(size)
    st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
        st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)

normal=doc.styles['Normal']; set_font(normal)
pf=normal.paragraph_format
pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
pf.first_line_indent=Cm(1.25); pf.line_spacing=1.5
pf.space_after=Pt(0); pf.widow_control=True
for name,size in [('Title',16),('Heading 1',16),('Heading 2',12),('Heading 3',12)]:
    st=doc.styles[name]; set_font(st,size,True)
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before=Pt(24 if name=='Heading 1' else 18 if name in ['Heading 2','Heading 3'] else 0)
    st.paragraph_format.space_after=Pt(18)
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 1'].paragraph_format.space_after=Pt(24)
for name in ['Caption','Source','Reference','Front','ListEntry','TOC Heading']:
    if name not in doc.styles:
        doc.styles.add_style(name,1)
    st=doc.styles[name];set_font(st,10 if name in ['Caption','Source'] else 12)
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.line_spacing=1.0
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Caption'].paragraph_format.keep_with_next=True
doc.styles['Source'].paragraph_format.space_after=Pt(12)
doc.styles['Reference'].paragraph_format.space_after=Pt(12)
doc.styles['Reference'].paragraph_format.keep_together=True
doc.styles['Front'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['TOC 1','TOC 2']:
    if name not in doc.styles: doc.styles.add_style(name,1)
    st=doc.styles[name];set_font(st,12,name=='TOC 1')
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.space_before=Pt(0)
    st.paragraph_format.space_after=Pt(0)

def para(text='', style=None):
    return doc.add_paragraph(text,style)

def front(text, before=0, after=0, bold=False, size=12):
    p=para(text,'Front');p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    for r in p.runs:r.bold=bold;r.font.size=Pt(size)
    return p

def page():doc.add_page_break()

def field(p, code, placeholder=''):
    run=p.add_run(); b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');run._r.append(b)
    r=p.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' '+code+' ';r._r.append(i)
    r=p.add_run();s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');r._r.append(s)
    p.add_run(placeholder)
    r=p.add_run();e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)

def section(start=None):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.header.is_linked_to_previous=False
    s.footer.is_linked_to_previous=False
    for n in s._sectPr.findall(qn('w:pgNumType')):
        s._sectPr.remove(n)
    if start is not None:
        n=OxmlElement('w:pgNumType');n.set(qn('w:start'),str(start));s._sectPr.append(n)
    return s

def unnumbered(title, include=False, newpage=True):
    if newpage:page()
    p=para(title,'Heading 1' if include else 'Title')
    p.paragraph_format.page_break_before=False
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p

front('ESCOLA ESTADUAL DEPUTADO DOMINGOS DE FIGUEIREDO',before=71,bold=True,size=14)
front('TURMA DO 3º ANO DE DESENVOLVIMENTO DE SISTEMAS',before=144,bold=True,size=14)
p=front('O SEGREDO DE VARGINHA',before=161,bold=True,size=18);p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0
p=front('Relatório do desenvolvimento do projeto',size=14)
for r in p.runs:r.italic=True
front('VARGINHA, 2026',before=228,bold=True,size=14)

face=section(1)
face.bottom_margin=Cm(2)
front('TURMA DO 3º ANO DE DESENVOLVIMENTO DE SISTEMAS',bold=True,size=14)
p=front('O SEGREDO DE VARGINHA',before=175,bold=True,size=18);p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.0
p=front('Relatório do desenvolvimento do projeto',size=14)
for r in p.runs:r.italic=True
p=para('Relatório do desenvolvimento do projeto apresentado ao 3º ano do curso técnico em Desenvolvimento de Sistemas da Escola Estadual Deputado Domingos de Figueiredo.','Front')
p.paragraph_format.left_indent=Cm(6.8);p.paragraph_format.space_before=Pt(93);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
p=para('Professores: Renan e Alex.','Front');p.paragraph_format.left_indent=Cm(6.8);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
front('VARGINHA, 2026',before=264,bold=True,size=14)


for heading,label,bookmark in [('LISTA DE ILUSTRAÇÕES','Figura 1 - Escola Industrial no acervo da remasterização','Figura1'),('LISTA DE TABELAS','Tabela 1 - Principais marcos do desenvolvimento','Tabela1')]:
    unnumbered(heading)
    p=para(label,'ListEntry');p.paragraph_format.line_spacing=1.5
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9),WD_ALIGN_PARAGRAPH.RIGHT,1)
    p.add_run('\t');field(p,f'PAGEREF {bookmark} \\h','0')
unnumbered('SUMÁRIO')
field(para('', 'ListEntry'),'TOC \\o "1-2" \\h \\z \\u','Atualizar sumário')
s=section();p=s.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p.paragraph_format.first_line_indent=Cm(0);field(p,'PAGE','0')
for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
book_id=0;fig_no=0
def caption(kind,n,title):
    global book_id
    p=para(f'{kind} {n} – {title}','Caption')
    book_id+=1
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(book_id));a.set(qn('w:name'),f'{kind}{n}')
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(book_id))
    p._p.insert(0,a);p._p.append(b)
    return p

def source(text):para('Fonte: '+text,'Source')

def table(kind,n,title,headers,rows,widths,src):
    caption(kind,n,title)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,w,h in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=h
    trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,w,text in zip(cells,widths,row):c.width=Cm(w);c.text=str(text)
    for rn,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        for cn,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side)
                draw=(kind=='Quadro') or (side=='top' and rn==0) or (side=='bottom' and (rn==0 or rn==len(t.rows)-1))
                e.set(qn('w:val'),'single' if draw else 'nil');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9' if kind=='Quadro' else '000000');borders.append(e)
            tcpr.append(borders)
            if rn==0 and kind=='Quadro':
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.keep_with_next=(kind=='Quadro')
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cn>0 or len(headers)==2 else WD_ALIGN_PARAGRAPH.CENTER
                if kind=='Tabela' and n==1 and cn==1:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=rn==0
    source(src)


content=(OUT/'conteudo.md').read_text(encoding='utf-8')
for chunk in re.split(r'\n\s*\n',content.strip()):
    if chunk.startswith('# '):
        p=para(chunk[2:],'Heading 1')
        if chunk.startswith('# 1 '):p.paragraph_format.page_break_before=False
    elif chunk.startswith('## '):para(chunk[3:],'Heading 2')
    elif chunk=='@page':page()
    elif chunk=='@timeline':
        table('Tabela',1,'Principais marcos do desenvolvimento',['Período','Marco documentado'],[
          ('Setembro','Protótipo, testes iniciais, entrada, sprites, combate e interações.'),
          ('1 a 3 de outubro','Expansão experimental, mapas, investigação e Livro do Tombo.'),
          ('4 a 7 de outubro','Etapas finais, chefe e condensação em 14 capítulos.'),
          ('8 a 21 de outubro','Remasterização, versões alpha e beta, abertura e sons.'),
          ('22 e 23 de outubro','Rotina, mesa escolar, falas, build 0.4.8 e nova revisão.')
        ],[3.4,12.6],'elaboração a partir do histórico e dos relatórios internos (Projeto O Segredo de Varginha, 2026a, 2026c).')
    elif chunk=='@fig':
        caption('Figura',1,'Escola Industrial no acervo da remasterização')
        p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'Docs/Remaster20261008/Prints/02_Escola_Industrial.png'),width=Cm(11))
        source('acervo interno Remaster20261008, de outubro de 2026.')
    else:para(chunk.replace('\n',' '))
unnumbered('5 REFERÊNCIAS BIBLIOGRÁFICAS',include=True)
refs=[
'CHACON, Scott; STRAUB, Ben. Pro Git. 2. ed. [S. l.]: Apress, 2014. Disponível em: https://git-scm.com/book/en/v2/Getting-Started-About-Version-Control. Acesso em: 23 out. 2026.',
'HUNICKE, Robin; LEBLANC, Marc; ZUBEK, Robert. MDA: a formal approach to game design and game research. In: AAAI WORKSHOP ON CHALLENGES IN GAME AI, 2004. Anais [...]. [S. l.]: AAAI, 2004. Disponível em: https://www.cs.northwestern.edu/~hunicke/MDA.pdf. Acesso em: 23 out. 2026.',
'MODELO ABTN. [S. l.: s. n.], [s. d.]. Arquivo PDF fornecido pelo solicitante como referência de apresentação.',
'PROJETO O SEGREDO DE VARGINHA. Histórico e código do projeto. Varginha, 2026a. Repositório local: 49 commits alcançáveis, de 8dd459a a a2946be, e alterações locais. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Documentos de design e campanha. Varginha, 2026b. Documentos internos. Localização: Docs/CampanhaOficial/Referencias/GDD_Usuario.md; O_Segredo_de_Varginha_LDD.md; Docs/CampanhaOficial/CAMPANHA_14_FASES.md. Consulta em: 23 out. 2026.',
'PROJETO O SEGREDO DE VARGINHA. Relatórios e evidências de desenvolvimento. Varginha, 2026c. Acervo interno em Docs e output, incluindo dialogue-casual-2026-10-23 e revisao-mestre-2026-10-23. Consulta em: 23 out. 2026.'
]
for ref in refs:para(ref,'Reference')
for st in doc.styles:
    if st.type==1:
        st.font.color.rgb=RGBColor(0,0,0)
        for b in st.element.xpath('./w:pPr/w:pBdr'):b.getparent().remove(b)
doc.core_properties.title='O Segredo de Varginha Relatório do desenvolvimento do projeto'
doc.core_properties.author='Turma do 3º ano de Desenvolvimento de Sistemas'
doc.core_properties.subject='Síntese documental até 23 de outubro de 2026'
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd)
dest=OUT/(NAME+'.docx');doc.save(dest)
print(json.dumps({'docx':str(dest),'words_body':len(content.split()),'sections':5,'references':len(refs)},ensure_ascii=False))
