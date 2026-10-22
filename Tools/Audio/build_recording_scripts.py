"""Build current-campaign recording scripts from exact speech and explicitly authored dialogue."""
import csv
import json
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from export_voice_script import ROOT, OUT, line_id

DELIVERY=ROOT/'output/dublagem-2026-10-22'
CAST={
    'Edelzio':'Adulto, 36 anos. Professor de Desenvolvimento de Sistemas. Cotidiano e raciocínio observável no início; tensão crescente, sem pânico constante. No final, compreensão e alívio.',
    'Edelzio criança':'A mesma pessoa aos seis anos em 1996. Curioso, frases simples, medo contido. Gravar limpo, sem imitar efeito de rádio. A frase Não deixa ela sair pertence ao surto do Fusca e à memória completa, nunca ao prólogo inicial.',
    'Renan':'Apoio técnico e professor já conhecido. Método, comparação de provas e cuidado com o grupo. Não falar como se dominasse a origem da interferência.',
    'Ouzana':'Bióloga. Observa e testa antes de concluir. Sua surpresa aparece quando a evidência contradiz a leitura normal. Não antecipar a função real da passagem.',
    'Padre Fábio':'Guardião dos registros. No começo, cautela e evasão. Na revelação, responsabilidade pelo ocultamento. Na luta, firmeza e proteção.',
    'Rodrigo':'Guia da interface. Fala com quem joga, sem presença física ou participação nos acontecimentos. Dicas opcionais devem orientar observação sem entregar a resposta.',
    'Apresentador do noticiário':'Noticiário fictício da abertura de 1996. Tom jornalístico sóbrio. Relatar alegações, sem confirmar nave, origem extraterrestre ou acontecimentos reais não fornecidos.'
}
STUDENTS=['Yasmin','Pedro','Matias','Fabio','Marcos','Anna Sabia','Ana Tavares','Luis Miguel Messias','Luis Martins']
for name in STUDENTS:
    CAST[name]='Aluno da Industrial. Conversa cotidiana nas interações opcionais; chamada curta e legível durante o apoio. As falas não atribuem uma personalidade ou poder novo.'
CAST['Fabio']+=' Este é o aluno Fabio, personagem distinto do Padre Fábio.'


def slug(text):
    return re.sub(r'[^a-z0-9]+','_',unicodedata.normalize('NFKD',text).encode('ascii','ignore').decode().lower()).strip('_')


def scene(row):
    source=row['origem'];speaker=row['personagem'];text=row['texto']
    if speaker in STUDENTS:return 'A07 A Industrial'
    if speaker=='Edelzio criança':return 'A06 O surto do Fusca'
    if speaker=='Rodrigo':return 'A08 Os fragmentos do mapa'
    if 'CampaignHouseSequence' in source:return 'A04 A chave na mesa do café'
    if 'CampaignExpansionController' in source:return 'A11 O laboratório de Ouzana' if speaker=='Ouzana' else 'A10 O arquivo de Fábio'
    if 'CampaignContinuation' in source:return 'A15 A página escondida' if speaker=='Padre Fábio' else 'A13 Arquivos e cópias preservadas'
    if 'CampaignStage' in source:
        return 'A07 A Industrial' if speaker=='Renan' or 'Renan' in text or '23:23' in text else 'A04 A chave na mesa do café'
    return 'A07 A Industrial'


def records():
    exact=json.loads((OUT/'falas_existentes.json').read_text(encoding='utf-8'))
    # The user requested only the present cast/scenes, not unused legacy flows or future NPCs/endings.
    active=[r for r in exact if r['escopo'] in ('CAMPANHA ATUAL','GDD E CUTSCENE ATUAL')]
    rows=[]
    for r in active:
        rows.append(dict(id=r['id'],personagem=r['personagem'],cena=scene(r),direcao=CAST[r['personagem']].split('.')[0]+'.',fala=r['texto'],estado='E',origem=r['origem']))
    with (OUT/'falas_propostas.tsv').open(encoding='utf-8') as file:
        for r in csv.DictReader(file,delimiter='\t'):
            rows.append(dict(id=line_id(r['personagem'],r['fala']),personagem=r['personagem'],cena=r['cena'],direcao=r['direcao'],fala=r['fala'],estado='E' if r['escopo']=='CITAÇÃO DO GDD PARA GRAVAÇÃO' and r['fala']=='Não deixa ela sair.' else 'P',origem='Redação para a cena indicada, baseada no GDD/LDD e na campanha atual'))
    unique={}
    for r in sorted(rows,key=lambda r:r['cena']):
        if r['id'] in unique:
            unique[r['id']]['direcao']+=' Uso adicional: '+r['cena']+'. '+r['direcao']
        else:unique[r['id']]=r
    rows=list(unique.values())
    # Supplemental vocal efforts are separate from story dialogue and never processed into fake speech.
    for i,(name,context,direction) in enumerate([
        ('Edelzio','Passo interrompido por susto','Um susto curto, sem palavra. Duas intensidades.'),
        ('Edelzio','Ignição quente','Reação curta à queimadura, sem prolongar o grito.'),
        ('Edelzio','Fuga','Três respirações de esforço curtas, com espaço para montagem.'),
        ('Edelzio','Queda e recuperação','Impacto vocal breve e recuperação do fôlego.'),
        ('Edelzio criança','Clarão','Reação assustada curta; não incluir a frase do Fusca.'),
        ('Padre Fábio','Proteção','Esforço breve ao sustentar a proteção.'),
        ('Ouzana','Arremesso do frasco','Expiração curta de esforço.'),
        ('Renan','Interferência inesperada','Susto contido e recuperação da fala.')]):
        rows.append(dict(id=f'esforco_{slug(name)}_{i+1:02}',personagem=name,cena='A22 Esforços vocais',direcao=direction,fala=f'[{context}: sem palavras.]',estado='V',origem='Direção de gravação, sem novo fato narrativo'))
    return rows


def introduction(ed=False):
    return [
        'Material para a sessão de dublagem de 22 de outubro de 2026. '+('Este arquivo contém somente Edelzio adulto e Edelzio criança.' if ed else 'Este arquivo reúne o elenco da campanha atual, em ordem de cena.'),
        'E indica texto existente no jogo ou citação já usada pela cutscene. P indica redação proposta para uma cena atual, ainda sujeita à revisão antes de integrar. V indica esforço vocal sem palavras. Se houver versões equivalentes E e P, escolher a versão aprovada e não tocar as duas na mesma cena.',
        'Gravar apenas o texto da fala. Nome, ID, intenção, contexto, colchetes, comandos e indicação de som não são falados. Separar as tomadas. O texto curto de coleta e das ações é aviso visual e não precisa de dublagem.',
        'A campanha atual determina a ordem das cenas. As fontes possuem numerações diferentes: não usar o número de fase do LDD para renomear os arquivos. Este roteiro usa 1898, 1996 e 2026 e o horário 23:23, conforme a versão jogável. Não antecipar o acordo ou a motivação da entidade antes da descoberta correspondente.',
        'Não há falas de professores ou NPCs ainda ausentes, nem blocos dos finais não jogáveis. A entidade comunica-se por interferências e memórias na campanha atual; não foi criado um monólogo humano para ela. Manifestações e criaturas usam os efeitos sonoros do jogo.',
        'Nome sugerido do áudio: ID_t01.wav e ID_t02.wav. Entregar voz seca, preferencialmente WAV mono a 48 kHz e 24 bits, com início e fim preservados e sem música, eco, glitch ou normalização agressiva. O som do rádio será aplicado na montagem. Os IDs E são compatíveis com o novo sistema; IDs P só serão vinculados depois da revisão e integração da respectiva fala.'
    ]


def markdown(rows,title,ed=False):
    lines=['# '+title,'']
    for p in introduction(ed):lines.extend([p,''])
    names=sorted({r['personagem'] for r in rows})
    lines.extend(['## Direção do elenco',''])
    for name in names:lines.extend(['### '+name,'',CAST[name],''])
    current=None
    for r in rows:
        if r['cena']!=current:current=r['cena'];lines.extend(['## '+current,''])
        lines.extend([f"**{r['personagem']} · {r['id']} · {r['estado']}**",'',r['direcao'],'',r['fala'],'',f"Fonte: {r['origem']}",''])
    lines.extend(['## Fontes','', '[GDD](../Referencias/GDD_Usuario.md) e [LDD](../Referencias/O_Segredo_de_Varginha_LDD.md). Falas existentes rastreadas em falas_existentes.json. As propostas apenas dramatizam informações já previstas para a cena, sem alterar o jogo por si só.',''])
    return '\n'.join(lines)


def word(rows,title,path,ed=False):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11.5);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.16
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(13)
    doc.add_paragraph(title,style='Title');doc.add_paragraph('O Segredo de Varginha  22 de outubro de 2026',style='Subtitle')
    for p in introduction(ed):doc.add_paragraph(p)
    doc.add_heading('Direção do elenco',1)
    for name in sorted({r['personagem'] for r in rows}):
        doc.add_heading(name,2);doc.add_paragraph(CAST[name])
    current=None
    for r in rows:
        if r['cena']!=current:current=r['cena'];doc.add_heading(current,1)
        meta=doc.add_paragraph();meta.paragraph_format.keep_with_next=True;meta.paragraph_format.space_before=Pt(9);meta.paragraph_format.space_after=Pt(3)
        run=meta.add_run(f"{r['personagem']}   {r['id']}   {r['estado']}");run.bold=True;run.font.size=Pt(9)
        cue=doc.add_paragraph(r['direcao']);cue.paragraph_format.keep_with_next=True;cue.paragraph_format.space_after=Pt(4)
        for run in cue.runs:run.italic=True;run.font.size=Pt(10)
        line=doc.add_paragraph(r['fala']);line.paragraph_format.keep_together=True;line.paragraph_format.space_after=Pt(10)
        for run in line.runs:run.font.size=Pt(12)
    # Source locators accompany each entry in Markdown/CSV; avoid an isolated final sources page.
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run('Dublagem 22 10 2026   ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(path)


def main():
    DELIVERY.mkdir(parents=True,exist_ok=True);rows=records()
    (OUT/'roteiro_gravacao.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
    with (OUT/'ROTEIRO_GRAVACAO.csv').open('w',encoding='utf-8-sig',newline='') as file:
        writer=csv.DictWriter(file,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    ed=[r for r in rows if r['personagem'] in ('Edelzio','Edelzio criança')]
    for name,title,subset,only in [('ROTEIRO_COMPLETO','Roteiro completo de dublagem',rows,False),('ROTEIRO_EDELZIO','Roteiro de dublagem do Edelzio',ed,True)]:
        text=markdown(subset,title,only);(OUT/(name+'.md')).write_text(text,encoding='utf-8');(DELIVERY/(name+'.md')).write_text(text,encoding='utf-8');word(subset,title,DELIVERY/(name+'.docx'),only)
    by=OUT/'por_personagem';by.mkdir(exist_ok=True)
    for name in sorted({r['personagem'] for r in rows}):
        selected=[r for r in rows if r['personagem']==name]
        (by/(slug(name)+'.md')).write_text(markdown(selected,'Falas de '+name,name.startswith('Edelzio')),encoding='utf-8')
    with zipfile.ZipFile(DELIVERY/'Roteiros_por_personagem.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(by.glob('*.md')):z.write(p,p.name)
        z.write(OUT/'ROTEIRO_GRAVACAO.csv','ROTEIRO_GRAVACAO.csv')
    summary=dict(total=len(rows),elenco=len({r['personagem'] for r in rows}),edelzio=len(ed),estados=dict(Counter(r['estado'] for r in rows)),personagens=dict(Counter(r['personagem'] for r in rows)))
    (DELIVERY/'manifesto.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(summary,ensure_ascii=False))


if __name__=='__main__':main()
